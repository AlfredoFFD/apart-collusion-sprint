"""Post-process the pandoc DOCX: grid borders on every table, 8.5 pt table text, wider "observed" column in the ladder table.
    python3 scripts/docx_polish.py report/report.docx
"""
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


def set_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0"); el.set(qn("w:color"), "999999")
        borders.append(el)
    for old in tblPr.findall(qn("w:tblBorders")):
        tblPr.remove(old)
    # schema order matters (Google Docs misreads an out-of-order tblPr): tblBorders goes before shd/tblLayout/tblCellMar/tblLook
    after = [c for c in tblPr if c.tag in (qn("w:shd"), qn("w:tblLayout"), qn("w:tblCellMar"), qn("w:tblLook"))]
    if after:
        after[0].addprevious(borders)
    else:
        tblPr.append(borders)


def set_col_widths(table, widths_in):
    for row in table.rows:
        for i, w in enumerate(widths_in):
            if i < len(row.cells):
                row.cells[i].width = Inches(w)


def main(path):
    d = Document(path)
    import os
    for t in d.tables:
        if os.environ.get("NO_BORDERS") != "1":
            set_borders(t)
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0); p.paragraph_format.space_after = Pt(0)
                    for r in p.runs:
                        r.font.size = Pt(8.5)
        # Google Docs merges a bordered table's header row into the next row when the row carries w:tblHeader; drop the marker
        for row in t.rows:
            trPr = row._tr.trPr
            if trPr is not None:
                for h in trPr.findall(qn("w:tblHeader")):
                    trPr.remove(h)
        header = " ".join(c.text.strip().lower() for c in t.rows[0].cells) if t.rows else ""
        if header.startswith("intervention (level)"):
            set_col_widths(t, [1.7, 0.9, 1.1, 3.3])
        elif header.startswith("runs detector"):
            set_col_widths(t, [1.9, 1.7, 1.4, 2.0])
        elif header.startswith("sellers market prompt"):
            set_col_widths(t, [0.9, 0.8, 2.6, 0.6, 2.1])
    d.save(path)
    print("polished", path, "tables:", len(d.tables))


if __name__ == "__main__":
    main(sys.argv[1])
